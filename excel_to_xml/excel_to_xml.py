#!/usr/bin/env python3
"""Convert an Excel worksheet to XML.

The first row of the sheet supplies the element names for each column.
Every row below it becomes one element with one child per column:

    <rows>
      <row>
        <name>Ada</name>
        <age>36</age>
      </row>
    </rows>
"""

import argparse
import re
import sys
import xml.etree.ElementTree as ET
from datetime import date, datetime, time

from openpyxl import load_workbook

INVALID_TAG_CHARS = re.compile(r"[^A-Za-z0-9_.-]+")


def to_tag(header):
    """Turn a header cell into a valid XML element name, or None if empty."""
    if header is None or not str(header).strip():
        return None
    tag = INVALID_TAG_CHARS.sub("_", str(header).strip()).strip("_") or "column"
    if not re.match(r"[A-Za-z_]", tag):
        tag = f"_{tag}"
    if tag.lower().startswith("xml"):
        tag = f"_{tag}"
    return tag


def unique_tags(headers):
    """Return a list of (column_index, tag) for usable columns, with duplicates renamed."""
    columns = []
    seen = set()
    for index, header in enumerate(headers):
        tag = to_tag(header)
        if tag is None:
            continue
        candidate, n = tag, 2
        while candidate in seen:
            candidate = f"{tag}_{n}"
            n += 1
        seen.add(candidate)
        columns.append((index, candidate))
    return columns


def to_text(value):
    """Format a cell value as XML text."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)


def sheet_to_element(ws, root_tag="rows", row_tag="row"):
    """Build an XML tree from a worksheet."""
    root = ET.Element(root_tag)
    rows = ws.iter_rows(values_only=True)

    header = next(rows, None)
    if header is None:
        return root
    columns = unique_tags(header)

    for values in rows:
        if all(v is None or not str(v).strip() for v in values):
            continue  # skip blank rows
        row = ET.SubElement(root, row_tag)
        for index, tag in columns:
            value = values[index] if index < len(values) else None
            ET.SubElement(row, tag).text = to_text(value)

    return root


def convert(input_path, sheet=None, root_tag="rows", row_tag="row"):
    """Return the XML tree for one sheet of an Excel file."""
    wb = load_workbook(input_path, read_only=True, data_only=True)
    try:
        if sheet is None:
            ws = wb.active
        elif sheet in wb.sheetnames:
            ws = wb[sheet]
        else:
            raise ValueError(f"Sheet {sheet!r} not found. Available: {', '.join(wb.sheetnames)}")
        root = sheet_to_element(ws, root_tag, row_tag)
    finally:
        wb.close()

    ET.indent(root, space="  ")
    return root


def main(argv=None):
    parser = argparse.ArgumentParser(description="Convert an Excel sheet to XML.")
    parser.add_argument("input", help="path to the .xlsx file")
    parser.add_argument("-o", "--output", help="output XML file (default: stdout)")
    parser.add_argument("-s", "--sheet", help="sheet name (default: active sheet)")
    parser.add_argument("--root", default="rows", help="root element name (default: rows)")
    parser.add_argument("--row", default="row", help="element name for each row (default: row)")
    args = parser.parse_args(argv)

    try:
        root = convert(args.input, args.sheet, args.root, args.row)
    except (OSError, ValueError) as err:
        parser.exit(1, f"error: {err}\n")

    tree = ET.ElementTree(root)
    if args.output:
        tree.write(args.output, encoding="utf-8", xml_declaration=True)
        print(f"Wrote {len(root)} rows to {args.output}", file=sys.stderr)
    else:
        sys.stdout.buffer.write(ET.tostring(root, encoding="utf-8", xml_declaration=True))
        sys.stdout.write("\n")


if __name__ == "__main__":
    main()
